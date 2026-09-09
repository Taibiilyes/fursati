# -*- coding: utf-8 -*-
"""توليد نموذج سيرة ذاتية عربي جاهز للتجربة (DOCX مع صورة شخصية)."""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = os.path.dirname(os.path.abspath(__file__))
GREEN = RGBColor(0x0E, 0x5C, 0x42)
DARK = RGBColor(0x21, 0x2A, 0x27)
GRAY = RGBColor(0x5A, 0x64, 0x60)


def rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p.alignment = align
    pPr = p._p.get_or_add_pPr()
    pPr.append(pPr.makeelement(qn("w:bidi"), {}))
    return p


def run(p, text, size=11, bold=False, color=DARK):
    r = p.add_run(text or "")
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "Segoe UI"
    rPr = r._r.get_or_add_rPr()
    rf = rPr.makeelement(qn("w:rFonts"), {}); rf.set(qn("w:cs"), "Segoe UI"); rPr.append(rf)
    rPr.append(rPr.makeelement(qn("w:szCs"), {qn("w:val"): str(int(size * 2))}))
    return r


def para(doc, text="", size=11, bold=False, color=DARK, align=WD_ALIGN_PARAGRAPH.RIGHT, sb=0, sa=4):
    p = doc.add_paragraph(); rtl(p, align)
    p.paragraph_format.space_before = Pt(sb); p.paragraph_format.space_after = Pt(sa)
    if text:
        run(p, text, size, bold, color)
    return p


def heading(doc, text):
    p = para(doc, text, size=13, bold=True, color=GREEN, sb=10, sa=4)
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn("w:pBdr"), {})
    pBdr.append(pPr.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "8",
                                                 qn("w:space"): "2", qn("w:color"): "0E5C42"}))
    pPr.append(pBdr)


doc = Document()
for s in doc.sections:
    s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(1.4), Cm(1.2), Cm(1.6), Cm(1.6)

from docx.enum.table import WD_TABLE_ALIGNMENT
tbl = doc.add_table(rows=1, cols=2)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl._tbl.tblPr.append(tbl._tbl.tblPr.makeelement(qn("w:bidiVisual"), {}))
tbl.columns[0].width = Cm(12.5); tbl.columns[1].width = Cm(5)

c0 = tbl.cell(0, 0)
p = c0.paragraphs[0]; rtl(p)
run(p, "أمين قاسمي", 20, True, GREEN)
p = c0.add_paragraph(); rtl(p); run(p, "مهندس دولة في الكهروميكانيك", 12, True)
p = c0.add_paragraph(); rtl(p); run(p, "📱 0555 12 34 56  •  ✉ amine.kasmi.dz@gmail.com  •  📍 سطيف، الجزائر", 9.5, False, GRAY)

c1 = tbl.cell(0, 1)
p = c1.paragraphs[0]; rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
p.add_run().add_picture(os.path.join(BASE, "static", "img", "sample_photo.jpg"), width=Cm(3.2))

heading(doc, "نبذة مهنية")
para(doc, "مهندس دولة في الكهروميكانيك، خبرة تفوق 5 سنوات في الصيانة الصناعية والمعدات الثقيلة، "
          "إتقان أنظمة التحكم الآلي PLC وسلامة المنشآت HSE، جاهز للانتقال والعمل في ولايات الجنوب الجزائري.", 10.5)

heading(doc, "المهارات")
para(doc, "صيانة صناعية، ميكانيك، كهرباء، أوتوماتيك، HSE، Office، لوجستيك، قيادة رخصة B، تحليل بيانات", 10.5)

heading(doc, "الخبرة المهنية")
para(doc, "مهندس صيانة | مصنع النور للصناعات الميكانيكية - سطيف | 2021-2026 | إدارة فريق صيانة وقائية وتصحيحية لخطوط الإنتاج", 10.5)
para(doc, "تقني صيانة | ورشة الطريق لتجهيزات الشاحنات - سطيف | 2019-2021 | صيانة ميكانيكية وكهربائية للشاحنات الثقيلة", 10.5)

heading(doc, "التكوين والشهادات")
para(doc, "مهندس دولة في الكهروميكانيك | جامعة عباس لغرور - خنشلة | 2019", 10.5)

heading(doc, "اللغات")
para(doc, "العربية: لغة أم  •  الفرنسية: جيد جدا  •  الإنجليزية: متوسط", 10.5)

heading(doc, "شهادات ودورات")
para(doc, "شهادة HSE أساسيات السلامة الصناعية، دبلوم Excel المتقدم، دورة التحكم الآلي Siemens PLC", 10.5)

os.makedirs(os.path.join(BASE, "static", "samples"), exist_ok=True)
out = os.path.join(BASE, "static", "samples", "cv_modele.docx")
doc.save(out)
print("OK →", out)
