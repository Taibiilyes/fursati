# -*- coding: utf-8 -*-
"""محرك تخصيص الملفات: يولّد سيرة ذاتية ورسالة تحفيزية بصيغة DOCX
مخصّصتين حسب المنصب والشركة المستهدَفَين."""

import io
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GREEN = RGBColor(0x0E, 0x5C, 0x42)
DARK = RGBColor(0x21, 0x2A, 0x27)
GRAY = RGBColor(0x5A, 0x64, 0x60)
GOLD = RGBColor(0xA8, 0x7D, 0x1D)

MONTHS_AR = ["جانفي", "فيفري", "مارس", "أفريل", "ماي", "جوان", "جويلية", "أوت",
             "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]


def _rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p.alignment = align
    pPr = p._p.get_or_add_pPr()
    bidi = pPr.makeelement(qn("w:bidi"), {})
    pPr.append(bidi)
    return p


def _run(p, text, size=11, bold=False, color=DARK, font="Segoe UI"):
    r = p.add_run(text or "")
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rf)
    rf.set(qn("w:cs"), font)
    szCs = rPr.makeelement(qn("w:szCs"), {qn("w:val"): str(int(size * 2))})
    rPr.append(szCs)
    if bold:
        bCs = rPr.makeelement(qn("w:bCs"), {})
        rPr.append(bCs)
    return r


def _para(doc, text="", size=11, bold=False, color=DARK, align=WD_ALIGN_PARAGRAPH.RIGHT,
          space_after=4, space_before=0):
    p = doc.add_paragraph()
    _rtl(p, align)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if text:
        _run(p, text, size, bold, color)
    return p


def _heading(doc, text):
    p = _para(doc, text, size=13, bold=True, color=GREEN, space_before=10, space_after=4)
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = pPr.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "8",
                                              qn("w:space"): "2", qn("w:color"): "0E5C42"})
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def _table_rtl(tbl):
    tblPr = tbl._tbl.tblPr
    tblPr.append(tblPr.makeelement(qn("w:bidiVisual"), {}))


def _today_ar():
    d = date.today()
    return f"{d.day} {MONTHS_AR[d.month - 1]} {d.year}"


def _years_experience(experience):
    import re
    total = 0.0
    for e in experience or []:
        p = e.get("period", "") or ""
        m = re.findall(r"(19|20)\d{2}", p)
        if "-" in p or "–" in p or "إلى" in p:
            yrs = [int(x) for x in re.findall(r"((?:19|20)\d{2})", p)]
            if len(yrs) >= 2:
                total += max(0, yrs[-1] - yrs[0])
        elif re.search(r"منذ|depuis|present|الآن", p, re.I) and m:
            total += max(0, date.today().year - int(m[0]))
    if total == 0 and experience:
        total = len(experience)
    return int(round(total))


# ══════════════════════ السيرة الذاتية المخصّصة ══════════════════════

def build_tailored_cv(profile, job, matched_skills):
    doc = Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(1.4), Cm(1.2)
        s.left_margin, s.right_margin = Cm(1.6), Cm(1.6)

    # الترويسة: جدول بعمودين (الاسم والمعلومات يميناً، الصورة يساراً)
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    _table_rtl(tbl)
    tbl.columns[0].width = Cm(12.5)
    tbl.columns[1].width = Cm(5.0)

    c_name = tbl.cell(0, 0)
    p = c_name.paragraphs[0]
    _rtl(p)
    _run(p, profile.get("full_name") or "مترشح", size=20, bold=True, color=GREEN)
    p2 = c_name.add_paragraph()
    _rtl(p2)
    _run(p2, profile.get("title") or "", size=12, bold=True, color=DARK)
    p3 = c_name.add_paragraph()
    _rtl(p3)
    contact = " • ".join(x for x in [profile.get("phone"), profile.get("email"),
                                     profile.get("wilaya")] if x)
    _run(p3, contact, size=9.5, color=GRAY)
    p4 = c_name.add_paragraph()
    _rtl(p4)
    _run(p4, f"مترشَّح لمنصب: {job['title']} — {job['company']}", size=11, bold=True, color=GOLD)

    c_photo = tbl.cell(0, 1)
    pp = c_photo.paragraphs[0]
    _rtl(pp, WD_ALIGN_PARAGRAPH.CENTER)
    photo_path = profile.get("_photo_path")
    if photo_path and __import__("os").path.exists(photo_path):
        try:
            pp.add_run().add_picture(photo_path, width=Cm(3.2))
        except Exception:
            pass

    # الملخص المهني المخصّص
    _heading(doc, "الملخص المهني")
    years = _years_experience(profile.get("experience"))
    years_txt = f"تفوق {years} سنوات في" if years else "في"
    matched_txt = "، ".join(matched_skills[:5]) if matched_skills else "مجال التخصص"
    summary = (f"مترشح ذو خبرة {years_txt} مجال «{profile.get('title') or matched_txt}»، أترشح لمنصب "
               f"{job['title']} لدى {job['company']} بـ{job['city']} ({job['wilaya']}). "
               f"أُتقن مجموعة من الكفاءات الأساسية المطلوبة لهذا المنصب أبرزها: {matched_txt}. "
               f"أسعى للمساهمة في أنشطة {job['company']} في قطاع {job['sector']}، "
               f"مع جاهزية كاملة للعمل بنظام {job['contract']} والانتقال للمنطقة.")
    _para(doc, summary, size=10.5, space_after=6)

    # المهارات (المطابقة أولاً)
    _heading(doc, "المهارات والكفاءات")
    user_skills = profile.get("skills") or []
    ordered = [s for s in user_skills if s in matched_skills] + \
              [s for s in user_skills if s not in matched_skills]
    if not ordered:
        ordered = matched_skills
    p = _para(doc, "", space_after=2)
    for i, s in enumerate(ordered):
        star = "★ " if s in matched_skills else "• "
        _run(p, star + s + ("   " if i < len(ordered) - 1 else ""), size=10.5,
             bold=bool(s in matched_skills), color=(GOLD if s in matched_skills else DARK))

    # الخبرة
    _heading(doc, "الخبرة المهنية")
    for e in profile.get("experience") or []:
        line = " | ".join(x for x in [e.get("role"), e.get("org"), e.get("period")] if x)
        _para(doc, "◆ " + line, size=10.5, bold=True, space_after=1)
        if e.get("desc"):
            _para(doc, "   " + e["desc"], size=10, color=GRAY, space_after=4)

    # التكوين
    _heading(doc, "التكوين والشهادات")
    for ed in profile.get("education") or []:
        line = " | ".join(x for x in [ed.get("degree"), ed.get("school"), ed.get("year")] if x)
        _para(doc, "◆ " + line, size=10.5)

    # اللغات
    _heading(doc, "اللغات")
    _para(doc, "   ".join(profile.get("languages") or []), size=10.5)

    # شهادات إضافية
    if profile.get("certifications"):
        _heading(doc, "شهادات ودورات تكوينية")
        _para(doc, "   ".join("• " + c for c in profile["certifications"]), size=10.5)

    # الاهتمامات
    if profile.get("interests"):
        _heading(doc, "الاهتمامات")
        _para(doc, "   ".join("• " + c for c in profile["interests"]), size=10.5)

    # تذييل
    _para(doc, f"ملف مخصَّص آلياً لمنصب {job['title']} — {job['company']}  |  {_today_ar()}",
          size=8.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ══════════════════════ الرسالة التحفيزية المخصّصة ══════════════════════

def build_letter(profile, job, matched_skills, lang="ar"):
    doc = Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(1.8), Cm(1.5)
        s.left_margin, s.right_margin = Cm(2.0), Cm(2.0)

    years = _years_experience(profile.get("experience"))
    years_txt = f"أكثر من {years} سنوات" if years else "خبرة ميدانية"
    matched_txt = "، ".join(matched_skills[:5]) if matched_skills else "كفاءاتي في مجال التخصص"

    if lang == "fr":
        name = profile.get("full_name") or "Candidat"
        _para(doc, name, size=13, bold=True, color=GREEN, align=WD_ALIGN_PARAGRAPH.LEFT)
        contact = " – ".join(x for x in [profile.get("phone"), profile.get("email"), profile.get("wilaya")] if x)
        _para(doc, contact, size=9.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.LEFT)
        _para(doc, f"À l'attention du Service Recrutement – {job['company']}", size=10.5, bold=True,
              align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=8)
        _para(doc, f"{job['city']}, le {_today_ar()}", size=10, color=GRAY, align=WD_ALIGN_PARAGRAPH.LEFT)
        _para(doc, f"Objet : Candidature au poste de {job['title']}", size=12, bold=True,
              color=GOLD, space_before=6, space_after=8)
        _para(doc, "Madame, Monsieur,", size=11, align=WD_ALIGN_PARAGRAPH.RIGHT)
        body1 = (f"Très intéressé par le poste de {job['title']} proposé par {job['company']} à {job['city']} "
                 f"({job['wilaya']}), je vous soumets ma candidature. Votre entreprise, acteur majeur du secteur "
                 f"« {job['sector']} » en Algérie, représente pour moi une opportunité concrète de mettre "
                 f"mes compétences au service de projets structurants.")
        body2 = (f"Fort de {years_txt} d'expérience et maîtrisant notamment {matched_txt}, "
                 f"je suis prêt à m'investir pleinement, en {job['contract']}, et à m'installer dans la région. "
                 f"Mon parcours, détaillé dans le CV ci-joint, correspond aux exigences du poste.")
        body3 = ("Je me tiens à votre disposition pour un entretien afin de vous exposer plus en détail ma motivation. "
                 "Dans l'attente de votre retour, je vous prie d'agréer, Madame, Monsieur, l'expression de mes salutations distinguées.")
        for b in (body1, body2, body3):
            p = _para(doc, b, size=10.5, space_after=8)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _para(doc, name, size=11, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=8)
    else:
        _para(doc, profile.get("full_name") or "مترشح", size=13, bold=True, color=GREEN)
        contact = " – ".join(x for x in [profile.get("phone"), profile.get("email"), profile.get("wilaya")] if x)
        _para(doc, contact, size=9.5, color=GRAY)
        _para(doc, f"إلى السيد(ة) مسؤول التوظيف — {job['company']}", size=11, bold=True,
              align=WD_ALIGN_PARAGRAPH.LEFT, space_before=8)
        _para(doc, f"البريد الإلكتروني للتوظيف: {job['hr_email']}", size=9.5, color=GRAY,
              align=WD_ALIGN_PARAGRAPH.LEFT)
        _para(doc, f"{job['city']}، في {_today_ar()}", size=10, color=GRAY, space_before=2)
        _para(doc, f"الموضوع: طلب ترشح لمنصب {job['title']}", size=12.5, bold=True,
              color=GOLD, space_before=6, space_after=8)
        _para(doc, "تحية طيبة وبعد،", size=11)
        body1 = (f"إنني بكل اهتمام أودّ التقدم لمنصب {job['title']} المعروض من طرف شركة {job['company']} "
                 f"بمدينة {job['city']} ({job['wilaya']}). وقد لفت انتباهي التزام شركتكم المرموقة في قطاع "
                 f"«{job['sector']}» بمشاريعها التنموية الكبرى، وهو ما يجسّد طموحي المهني وأرغبت في أن أكون جزءاً منه.")
        body2 = (f"يتيح لي مسار المهني، الذي يمتد لـ{years_txt} في مجال تخصصي، إتقان كفاءات تتوافق مباشرة مع متطلبات "
                 f"المنصب، أبرزها: {matched_txt}. كما أنّ جاهزيتي الكاملة للعمل بنظام {job['contract']} "
                 f"والاستقرار بمنطقة {job['wilaya']} تعززان قدرتي على الاندماج الفوري في فرق العمل لديكم.")
        body3 = ("وقد أرفقت لكم سيرتي الذاتية المعدّة خصيصاً لهذا المنصب، وأبقى رهن إشارتكم لإجراء مقابلة شخصية "
                 "أقدّم خلالها دافعياتي ومؤهلاتي عن قرب. وتفضلوا بقبول فائق عبارات الاحترام والتقدير.")
        for b in (body1, body2, body3):
            p = _para(doc, b, size=10.5, space_after=8)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _para(doc, profile.get("full_name") or "مترشح", size=11, bold=True,
              align=WD_ALIGN_PARAGRAPH.LEFT, space_before=8)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
